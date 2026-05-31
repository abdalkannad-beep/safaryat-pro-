import zipfile
import os

from flask import Flask, render_template, request, redirect, url_for, session, flash, send_file, send_from_directory
import sqlite3
import threading
import time
import requests
import uuid
from pathlib import Path
import shutil
from datetime import datetime, date, timedelta
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from urllib.parse import quote
from io import BytesIO

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = "static/uploads/tickets"
app.secret_key = "CHANGE_THIS_SECRET_KEY_IN_PRODUCTION"
DB = "agency.db"

DEFAULT_CATEGORIES = [
    ("عمرة", 500),
    ("حج", 3000),
    ("عمل", 1500),
    ("زيارة", 1000),
    ("زيارة عمل", 1200),
    ("تعديل مهنة", 700),
    ("رخصة قيادة", 600),
]

def db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn

def log_activity(action, details=""):
    try:
        conn = db()
        conn.execute("""CREATE TABLE IF NOT EXISTS activity_logs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER,
            action TEXT NOT NULL,
            details TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""")
        conn.execute("INSERT INTO activity_logs(user_id, action, details) VALUES(?,?,?)",
                     (session.get("user_id"), action, details))
        conn.commit()
        conn.close()
    except Exception:
        pass


def column_exists(conn, table_name, column_name):
    try:
        cols = conn.execute(f"PRAGMA table_info({table_name})").fetchall()
        return any(c["name"] == column_name for c in cols)
    except Exception:
        return False

def ensure_database_upgrade():
    conn = db()
    c = conn.cursor()

    # ترقية جدول المعاملات للنسخ القديمة
    tx_columns = {
        "process_status": "TEXT DEFAULT 'ready'",
        "agent_name": "TEXT",
        "currency_code": "TEXT DEFAULT 'YER'",
        "exchange_rate": "REAL DEFAULT 1",
        "total_base": "REAL DEFAULT 0",
        "paid_base": "REAL DEFAULT 0",
        "remaining_base": "REAL DEFAULT 0",
        "is_deleted": "INTEGER DEFAULT 0",
        "deleted_at": "TEXT",
        "deleted_by": "INTEGER"
    }
    for col, typ in tx_columns.items():
        if not column_exists(conn, "transactions", col):
            try:
                c.execute(f"ALTER TABLE transactions ADD COLUMN {col} {typ}")
            except Exception:
                pass

    # ترقية جدول العملاء
    client_columns = {
        "is_deleted": "INTEGER DEFAULT 0",
        "deleted_at": "TEXT",
        "deleted_by": "INTEGER"
    }
    for col, typ in client_columns.items():
        if not column_exists(conn, "clients", col):
            try:
                c.execute(f"ALTER TABLE clients ADD COLUMN {col} {typ}")
            except Exception:
                pass

    # إنشاء جدول العملات
    c.execute("""CREATE TABLE IF NOT EXISTS currencies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        rate_to_base REAL NOT NULL DEFAULT 1,
        is_base INTEGER DEFAULT 0,
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    default_currencies = [
        ("YER", "ريال يمني", 1, 1),
        ("SAR", "ريال سعودي", 140, 0),
        ("USD", "دولار أمريكي", 530, 0),
        ("AED", "درهم إماراتي", 145, 0)
    ]
    for code, name, rate, is_base in default_currencies:
        if not c.execute("SELECT id FROM currencies WHERE code=?", (code,)).fetchone():
            c.execute("INSERT INTO currencies(code,name,rate_to_base,is_base) VALUES(?,?,?,?)", (code, name, rate, is_base))

    # إنشاء جدول المحذوفات
    c.execute("""CREATE TABLE IF NOT EXISTS deleted_records (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        table_name TEXT NOT NULL,
        record_id INTEGER NOT NULL,
        title TEXT,
        deleted_by INTEGER,
        deleted_at TEXT DEFAULT CURRENT_TIMESTAMP,
        restored INTEGER DEFAULT 0,
        restored_at TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS activity_logs (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        action TEXT NOT NULL,
        details TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    conn.commit()
    conn.close()

@app.before_request
def auto_upgrade_database():
    try:
        ensure_database_upgrade()
    except Exception:
        pass


AR_DAYS = {
    "Saturday": "السبت",
    "Sunday": "الأحد",
    "Monday": "الإثنين",
    "Tuesday": "الثلاثاء",
    "Wednesday": "الأربعاء",
    "Thursday": "الخميس",
    "Friday": "الجمعة",
}

def arabic_day(dt=None):
    dt = dt or datetime.now()
    return AR_DAYS.get(dt.strftime("%A"), dt.strftime("%A"))


def init_db():
    conn = db()
    c = conn.cursor()

    c.execute("""CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        username TEXT UNIQUE NOT NULL,
        password_hash TEXT NOT NULL,
        full_name TEXT NOT NULL,
        role TEXT NOT NULL DEFAULT 'employee',
        can_edit INTEGER DEFAULT 1,
        can_delete INTEGER DEFAULT 0,
        is_active INTEGER DEFAULT 1,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS settings (
        key TEXT PRIMARY KEY,
        value TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS categories (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE NOT NULL,
        price REAL NOT NULL DEFAULT 0,
        active INTEGER DEFAULT 1,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS clients (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        full_name TEXT NOT NULL,
        passport_number TEXT,
        phone TEXT,
        address TEXT,
        notes TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS agencies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        owner_user_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS transactions (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        client_id INTEGER NOT NULL,
        category_id INTEGER NOT NULL,
        service_name TEXT NOT NULL,
        delegated TEXT DEFAULT 'لا',
        needs_license TEXT DEFAULT 'لا',
        needs_profession_change TEXT DEFAULT 'لا',
        status TEXT DEFAULT 'جديدة',
        process_status TEXT DEFAULT 'ready',
        start_date TEXT NOT NULL,
        alert_days INTEGER,
        total_amount REAL NOT NULL DEFAULT 0,
        paid_amount REAL NOT NULL DEFAULT 0,
        remaining_amount REAL NOT NULL DEFAULT 0,
        payment_status TEXT DEFAULT 'غير مسدد',
        payment_date TEXT,
        employee_id INTEGER,
        notes TEXT,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(client_id) REFERENCES clients(id),
        FOREIGN KEY(category_id) REFERENCES categories(id),
        FOREIGN KEY(employee_id) REFERENCES users(id)
    )""")

    # إضافة أعمدة للنسخ القديمة
    try:
        c.execute("ALTER TABLE transactions ADD COLUMN process_status TEXT DEFAULT 'ready'")
    except Exception:
        pass

    c.execute("""CREATE TABLE IF NOT EXISTS tickets (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        client_id INTEGER,
        ticket_type TEXT NOT NULL,
        route_from TEXT,
        route_to TEXT,
        travel_date TEXT,
        travel_time TEXT,
        seat_number TEXT,
        company_name TEXT,
        price REAL DEFAULT 0,
        paid_amount REAL DEFAULT 0,
        remaining_amount REAL DEFAULT 0,
        image_path TEXT,
        notes TEXT,
        created_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(client_id) REFERENCES clients(id)
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS payments (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        transaction_id INTEGER NOT NULL,
        amount REAL NOT NULL,
        paid_at TEXT NOT NULL,
        method TEXT DEFAULT 'نقد',
        notes TEXT,
        user_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY(transaction_id) REFERENCES transactions(id)
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS expenses (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        amount REAL NOT NULL,
        expense_date TEXT NOT NULL,
        notes TEXT,
        user_id INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS whatsapp_messages (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        client_id INTEGER,
        transaction_id INTEGER,
        phone TEXT NOT NULL,
        message TEXT NOT NULL,
        send_type TEXT DEFAULT 'manual',
        scheduled_at TEXT,
        status TEXT DEFAULT 'pending',
        sent_at TEXT,
        created_by INTEGER,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP,
        provider_response TEXT
    )""")

    c.execute("""CREATE TABLE IF NOT EXISTS notifications (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        title TEXT NOT NULL,
        body TEXT,
        type TEXT DEFAULT 'info',
        is_read INTEGER DEFAULT 0,
        created_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")

    super_admin = c.execute("SELECT id FROM users WHERE username='superadmin'").fetchone()
    if not super_admin:
        c.execute("INSERT INTO users (username,password_hash,full_name,role,can_edit,can_delete) VALUES (?,?,?,?,?,?)",
                  ("superadmin", generate_password_hash("super123"), "المشرف الأعلى", "super_admin", 1, 1))

    admin = c.execute("SELECT id FROM users WHERE username='admin'").fetchone()
    if not admin:
        c.execute("INSERT INTO users (username,password_hash,full_name,role,can_edit,can_delete) VALUES (?,?,?,?,?,?)",
                  ("admin", generate_password_hash("admin123"), "مدير الوكالة", "admin", 1, 1))

    if not c.execute("SELECT id FROM agencies LIMIT 1").fetchone():
        owner = c.execute("SELECT id FROM users WHERE username='admin'").fetchone()
        c.execute("INSERT INTO agencies(name, owner_user_id) VALUES(?,?)", ("سفريات برو", owner["id"] if owner else None))

    for k, v in {
        "default_alert_days":"",
        "agency_name":"سفريات برو",
        "whatsapp_api_url":"",
        "whatsapp_api_method":"GET",
        "whatsapp_api_headers":"",
        "whatsapp_enabled":"0",
        "manager_notify_phone":""
    }.items():
        if not c.execute("SELECT key FROM settings WHERE key=?", (k,)).fetchone():
            c.execute("INSERT INTO settings (key,value) VALUES (?,?)", (k, v))

    for name, price in DEFAULT_CATEGORIES:
        if not c.execute("SELECT id FROM categories WHERE name=?", (name,)).fetchone():
            c.execute("INSERT INTO categories (name,price) VALUES (?,?)", (name, price))

    conn.commit()
    conn.close()

def current_user():
    if "user_id" not in session:
        return None
    conn = db()
    user = conn.execute("SELECT * FROM users WHERE id=?", (session["user_id"],)).fetchone()
    conn.close()
    return user

def require_login():
    if not current_user():
        return redirect(url_for("login"))
    return None

def is_super_admin():
    u = current_user()
    return u and u["role"] == "super_admin"

def is_admin():
    u = current_user()
    return u and u["role"] in ["admin", "super_admin"]

def admin_required():
    if not is_admin():
        flash("هذه الصفحة خاصة بالمدير أو المشرف الأعلى فقط")
        return redirect(url_for("dashboard"))
    return None

def super_admin_required():
    if not is_super_admin():
        flash("هذه الصفحة خاصة بحساب المشرف الأعلى فقط")
        return redirect(url_for("dashboard"))
    return None


def get_setting(key, default=""):
    conn = db()
    row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return row["value"] if row else default

def set_notification(title, body="", ntype="info"):
    conn = db()
    conn.execute("INSERT INTO notifications(title, body, type) VALUES(?,?,?)", (title, body, ntype))
    conn.commit()
    conn.close()

def render_whatsapp_url(template, phone, message):
    return (template or "").replace("{phone}", phone or "").replace("{message}", quote(message or ""))

def send_whatsapp_api(phone, message):
    api_url = get_setting("whatsapp_api_url", "")
    method = get_setting("whatsapp_api_method", "GET").upper()
    enabled = get_setting("whatsapp_enabled", "0")

    if enabled != "1" or not api_url:
        return False, "إرسال واتساب التلقائي غير مفعل أو لم يتم إدخال رابط API"

    final_url = render_whatsapp_url(api_url, phone, message)

    try:
        if method == "POST":
            res = requests.post(final_url, json={"phone": phone, "message": message}, timeout=20)
        else:
            res = requests.get(final_url, timeout=20)

        ok = 200 <= res.status_code < 300
        response_text = (res.text or "")[:500]
        return ok, f"HTTP {res.status_code}: {response_text}"
    except Exception as e:
        return False, str(e)

def whatsapp_scheduler_loop():
    while True:
        try:
            conn = db()
            due = conn.execute("""SELECT * FROM whatsapp_messages
                                  WHERE send_type='scheduled'
                                  AND status='pending'
                                  AND scheduled_at IS NOT NULL
                                  AND datetime(scheduled_at) <= datetime('now', 'localtime')
                                  ORDER BY id ASC LIMIT 10""").fetchall()
            conn.close()

            for msg in due:
                ok, response = send_whatsapp_api(msg["phone"], msg["message"])
                conn = db()
                if ok:
                    conn.execute("""UPDATE whatsapp_messages
                                    SET status='sent', sent_at=CURRENT_TIMESTAMP, provider_response=?
                                    WHERE id=?""", (response, msg["id"]))
                    conn.execute("""INSERT INTO notifications(title, body, type)
                                    VALUES(?,?,?)""",
                                 ("تم إرسال رسالة واتساب", f"تم إرسال رسالة إلى {msg['phone']}", "success"))
                else:
                    conn.execute("""UPDATE whatsapp_messages
                                    SET status='failed', provider_response=?
                                    WHERE id=?""", (response, msg["id"]))
                    conn.execute("""INSERT INTO notifications(title, body, type)
                                    VALUES(?,?,?)""",
                                 ("فشل إرسال رسالة واتساب", f"الرقم: {msg['phone']} — السبب: {response}", "danger"))
                conn.commit()
                conn.close()
        except Exception:
            pass
        time.sleep(30)

_scheduler_started = False
def start_scheduler_once():
    global _scheduler_started
    if not _scheduler_started:
        t = threading.Thread(target=whatsapp_scheduler_loop, daemon=True)
        t.start()
        _scheduler_started = True


def can_manage_user(target):
    u = current_user()
    if not u or not target:
        return False
    if u["role"] == "super_admin":
        return True
    if u["role"] == "admin" and target["role"] == "employee":
        return True
    return False

@app.context_processor
def inject():
    unread_notifications = 0
    if current_user():
        conn = db()
        unread_notifications = conn.execute("SELECT COUNT(*) c FROM notifications WHERE is_read=0").fetchone()["c"]
        conn.close()
    return {"user": current_user(), "is_admin": is_admin, "is_super_admin": is_super_admin, "arabic_day": arabic_day, "unread_notifications": unread_notifications}

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        conn = db()
        user = conn.execute("SELECT * FROM users WHERE username=? AND is_active=1", (username,)).fetchone()
        conn.close()
        if user and check_password_hash(user["password_hash"], password):
            session["user_id"] = user["id"]
            return redirect(url_for("dashboard"))
        flash("بيانات الدخول غير صحيحة")
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/")
def dashboard():
    gate = require_login()
    if gate: return gate
    conn = db()
    today = date.today().isoformat()
    counts = {}
    counts["today_transactions"] = conn.execute("SELECT COUNT(*) c FROM transactions WHERE date(created_at)=date(?)", (today,)).fetchone()["c"]
    counts["clients"] = conn.execute("SELECT COUNT(*) c FROM clients").fetchone()["c"]
    counts["unpaid"] = conn.execute("SELECT COUNT(*) c FROM transactions WHERE remaining_amount>0").fetchone()["c"]
    counts["today_income"] = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE date(paid_at)=date(?)", (today,)).fetchone()["s"]
    counts["today_expenses"] = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE date(expense_date)=date(?)", (today,)).fetchone()["s"]

    rows = conn.execute("""SELECT t.*, c.full_name, c.passport_number, c.phone
        FROM transactions t JOIN clients c ON c.id=t.client_id ORDER BY t.id DESC LIMIT 12""").fetchall()
    alerts = []
    for r in rows:
        try:
            passed = (date.today() - datetime.strptime(r["start_date"], "%Y-%m-%d").date()).days
            if r["alert_days"] and passed >= int(r["alert_days"]):
                alerts.append({"name": r["full_name"], "service": r["service_name"], "days": passed, "limit": r["alert_days"]})
        except Exception:
            pass
    conn.close()
    return render_template("dashboard.html", counts=counts, rows=rows, alerts=alerts)

@app.route("/clients")
def clients():
    gate = require_login()
    if gate: return gate
    q = request.args.get("q","").strip()
    conn = db()
    if q:
        rows = conn.execute("""SELECT * FROM clients WHERE full_name LIKE ? OR passport_number LIKE ? OR phone LIKE ?
                               ORDER BY id DESC""", (f"%{q}%", f"%{q}%", f"%{q}%")).fetchall()
    else:
        rows = conn.execute("SELECT * FROM clients ORDER BY id DESC").fetchall()
    conn.close()
    return render_template("clients.html", rows=rows, q=q)

@app.route("/clients/add", methods=["GET","POST"])
def client_add():
    gate = require_login()
    if gate: return gate
    if request.method == "POST":
        conn = db()
        conn.execute("""INSERT INTO clients(full_name,passport_number,phone,address,notes)
                        VALUES(?,?,?,?,?)""",
                     (request.form["full_name"], request.form.get("passport_number"), request.form.get("phone"),
                      request.form.get("address"), request.form.get("notes")))
        conn.commit(); conn.close()
        flash("تم حفظ العميل")
        return redirect(url_for("clients"))
    return render_template("client_form.html", row=None)

@app.route("/transactions/add", methods=["GET","POST"])
def transaction_add():
    gate = require_login()
    if gate: return gate
    conn = db()
    cats = conn.execute("SELECT * FROM categories WHERE active=1 ORDER BY name").fetchall()
    clients_rows = conn.execute("SELECT * FROM clients ORDER BY full_name").fetchall()
    default_alert = conn.execute("SELECT value FROM settings WHERE key='default_alert_days'").fetchone()["value"]
    if request.method == "POST":
        client_id = request.form.get("client_id")
        if client_id == "new":
            cur = conn.execute("""INSERT INTO clients(full_name,passport_number,phone,address,notes)
                                  VALUES(?,?,?,?,?)""",
                               (request.form["new_full_name"], request.form.get("new_passport_number"),
                                request.form.get("new_phone"), "", ""))
            client_id = cur.lastrowid
        cat = conn.execute("SELECT * FROM categories WHERE id=?", (request.form["category_id"],)).fetchone()
        total = float(request.form.get("total_amount") or cat["price"] or 0)
        paid = float(request.form.get("paid_amount") or 0)
        remaining = total - paid
        payment_status = "مسدد" if remaining <= 0 else "غير مسدد"
        cur = conn.execute("""INSERT INTO transactions(
            client_id, category_id, service_name, delegated, needs_license, needs_profession_change, status, process_status,
            start_date, alert_days, total_amount, paid_amount, remaining_amount, payment_status, payment_date,
            employee_id, notes
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (client_id, cat["id"], cat["name"], request.form.get("delegated"), request.form.get("needs_license"),
         request.form.get("needs_profession_change"), request.form.get("status"), request.form.get("process_status", "ready"), request.form.get("start_date"),
         int(request.form.get("alert_days")) if request.form.get("alert_days") else None, total, paid, remaining, payment_status,
         request.form.get("payment_date"), session["user_id"], request.form.get("notes")))
        tx_id = cur.lastrowid
        if paid > 0:
            conn.execute("INSERT INTO payments(transaction_id,amount,paid_at,method,user_id,notes) VALUES(?,?,?,?,?,?)",
                         (tx_id, paid, request.form.get("payment_date") or date.today().isoformat(), "نقد", session["user_id"], "دفعة أولى"))
        conn.commit(); conn.close()
        log_activity("إضافة معاملة", f"تمت إضافة معاملة للعميل رقم {client_id}")
        flash("تم حفظ المعاملة")
        return redirect(url_for("transactions"))
    conn.close()
    return render_template("transaction_form.html", cats=cats, clients=clients_rows, default_alert=default_alert, today=date.today().isoformat())


@app.route("/transactions/batch", methods=["GET","POST"])
def transaction_batch():
    gate = require_login()
    if gate: return gate

    conn = db()
    cats = conn.execute("SELECT * FROM categories WHERE active=1 ORDER BY name").fetchall()
    default_alert = conn.execute("SELECT value FROM settings WHERE key='default_alert_days'").fetchone()["value"]

    if request.method == "POST":
        count = int(request.form.get("row_count", 0))
        saved = 0

        for i in range(count):
            full_name = request.form.get(f"full_name_{i}", "").strip()
            passport_number = request.form.get(f"passport_number_{i}", "").strip()
            phone = request.form.get(f"phone_{i}", "").strip()
            category_id = request.form.get(f"category_id_{i}", "").strip()

            if not full_name and not passport_number:
                continue

            cat = conn.execute("SELECT * FROM categories WHERE id=?", (category_id,)).fetchone()
            if not cat:
                continue

            cur_client = conn.execute("""INSERT INTO clients(full_name,passport_number,phone,address,notes)
                                         VALUES(?,?,?,?,?)""",
                                      (full_name or "بدون اسم", passport_number, phone, "", "إدخال جماعي"))
            client_id = cur_client.lastrowid

            total = float(request.form.get(f"total_amount_{i}") or cat["price"] or 0)
            paid = float(request.form.get(f"paid_amount_{i}") or 0)
            remaining = total - paid
            payment_status = "مسدد" if remaining <= 0 else "غير مسدد"
            now_date = request.form.get("entry_date") or date.today().isoformat()

            cur_tx = conn.execute("""INSERT INTO transactions(
                client_id, category_id, service_name, delegated, needs_license, needs_profession_change, status,
                start_date, alert_days, total_amount, paid_amount, remaining_amount, payment_status, payment_date,
                employee_id, notes
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (client_id, cat["id"], cat["name"], request.form.get(f"delegated_{i}", "لا"),
             request.form.get(f"needs_license_{i}", "لا"), request.form.get(f"needs_profession_change_{i}", "لا"),
             "جديدة", request.form.get(f"process_status_{i}", "ready"), now_date, int(request.form.get(f"alert_days_{i}")) if request.form.get(f"alert_days_{i}") else None,
             total, paid, remaining, payment_status, now_date if paid > 0 else "",
             session["user_id"], request.form.get(f"notes_{i}", "")))

            if paid > 0:
                conn.execute("INSERT INTO payments(transaction_id,amount,paid_at,method,user_id,notes) VALUES(?,?,?,?,?,?)",
                             (cur_tx.lastrowid, paid, now_date, "نقد", session["user_id"], "دفعة من الإدخال الجماعي"))
            saved += 1

        conn.commit()
        conn.close()
        flash(f"تم الحفظ النهائي لعدد {saved} معاملة")
        return redirect(url_for("transactions", day=date.today().isoformat()))

    now = datetime.now()
    conn.close()
    return render_template("transaction_batch.html", cats=cats, default_alert=default_alert,
                           today=date.today().isoformat(), now_time=now.strftime("%H:%M:%S"),
                           day_name=arabic_day(now))

@app.route("/transactions")
def transactions():
    gate = require_login()
    if gate: return gate
    q = request.args.get("q","").strip()
    day = request.args.get("day","").strip()
    process_filter = request.args.get("process_status","").strip()
    conn = db()
    sql = """SELECT t.*, c.full_name, c.passport_number, c.phone, u.full_name employee
             FROM transactions t
             JOIN clients c ON c.id=t.client_id
             LEFT JOIN users u ON u.id=t.employee_id
             WHERE 1=1"""
    params = []
    if q:
        sql += " AND (c.full_name LIKE ? OR c.passport_number LIKE ? OR c.phone LIKE ? OR t.service_name LIKE ?)"
        params += [f"%{q}%", f"%{q}%", f"%{q}%", f"%{q}%"]
    if day:
        sql += " AND date(t.created_at)=date(?)"
        params.append(day)
    if process_filter:
        sql += " AND t.process_status=?"
        params.append(process_filter)
    sql += " ORDER BY t.id DESC"
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return render_template("transactions.html", rows=rows, q=q, day=day, today=date.today().isoformat())

@app.route("/transactions/<int:tx_id>/pay", methods=["POST"])
def add_payment(tx_id):
    gate = require_login()
    if gate: return gate
    amount = float(request.form.get("amount") or 0)
    paid_at = request.form.get("paid_at") or date.today().isoformat()
    conn = db()
    tx = conn.execute("SELECT * FROM transactions WHERE id=?", (tx_id,)).fetchone()
    new_paid = float(tx["paid_amount"]) + amount
    new_remaining = float(tx["total_amount"]) - new_paid
    status = "مسدد" if new_remaining <= 0 else "غير مسدد"
    conn.execute("INSERT INTO payments(transaction_id,amount,paid_at,method,notes,user_id) VALUES(?,?,?,?,?,?)",
                 (tx_id, amount, paid_at, request.form.get("method","نقد"), request.form.get("notes",""), session["user_id"]))
    conn.execute("UPDATE transactions SET paid_amount=?, remaining_amount=?, payment_status=?, payment_date=? WHERE id=?",
                 (new_paid, new_remaining, status, paid_at, tx_id))
    conn.commit(); conn.close()
    flash("تمت إضافة الدفعة")
    return redirect(url_for("transactions"))

@app.route("/categories", methods=["GET","POST"])
def categories():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm
    conn = db()
    if request.method == "POST":
        conn.execute("INSERT OR IGNORE INTO categories(name,price,active) VALUES(?,?,1)",
                     (request.form["name"], float(request.form.get("price") or 0)))
        conn.commit()
    rows = conn.execute("SELECT * FROM categories ORDER BY active DESC, name").fetchall()
    conn.close()
    return render_template("categories.html", rows=rows)

@app.route("/categories/<int:cid>/update", methods=["POST"])
def category_update(cid):
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm
    conn = db()
    conn.execute("UPDATE categories SET name=?, price=?, active=? WHERE id=?",
                 (request.form["name"], float(request.form.get("price") or 0), int(request.form.get("active",1)), cid))
    conn.commit(); conn.close()
    return redirect(url_for("categories"))

@app.route("/users", methods=["GET","POST"])
def users():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    u = current_user()
    conn = db()

    if request.method == "POST":
        requested_role = request.form["role"]

        # المشرف الأعلى فقط يستطيع إنشاء أدمن أو مشرف أعلى آخر.
        if u["role"] == "admin" and requested_role != "employee":
            conn.close()
            flash("الأدمن يستطيع إنشاء موظفين فقط. إنشاء أدمن جديد خاص بالمشرف الأعلى.")
            return redirect(url_for("users"))

        # لا يسمح للأدمن بإنشاء حسابات حذف عالية.
        can_delete = int(request.form.get("can_delete", 0))
        if u["role"] == "admin":
            can_delete = 0

        conn.execute("""INSERT INTO users(username,password_hash,full_name,role,can_edit,can_delete,is_active)
                        VALUES(?,?,?,?,?,?,?)""",
                     (request.form["username"], generate_password_hash(request.form["password"]),
                      request.form["full_name"], requested_role, int(request.form.get("can_edit",1)),
                      can_delete, 1))
        conn.commit()
        flash("تم إنشاء الحساب")

    if u["role"] == "super_admin":
        rows = conn.execute("SELECT * FROM users ORDER BY id DESC").fetchall()
        allowed_roles = ["employee", "admin", "super_admin"]
    else:
        rows = conn.execute("SELECT * FROM users WHERE role='employee' OR id=? ORDER BY id DESC", (u["id"],)).fetchall()
        allowed_roles = ["employee"]

    conn.close()
    return render_template("users.html", rows=rows, allowed_roles=allowed_roles)


@app.route("/users/<int:uid>/status", methods=["POST"])
def user_status(uid):
    gate = require_login()
    if gate: return gate
    u = current_user()
    conn = db()
    target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not target:
        conn.close()
        flash("الحساب غير موجود")
        return redirect(url_for("users"))

    if u["role"] != "super_admin" and target["role"] != "employee":
        conn.close()
        flash("لا يمكنك التحكم بهذا الحساب")
        return redirect(url_for("users"))

    if target["role"] == "super_admin" and target["id"] == u["id"]:
        conn.close()
        flash("لا يمكن تعطيل حسابك الأعلى من نفس الصفحة")
        return redirect(url_for("users"))

    conn.execute("UPDATE users SET is_active=?, can_edit=?, can_delete=? WHERE id=?",
                 (int(request.form.get("is_active", 1)), int(request.form.get("can_edit", 1)),
                  int(request.form.get("can_delete", 0)), uid))
    conn.commit()
    conn.close()
    flash("تم تحديث صلاحيات الحساب")
    return redirect(url_for("users"))


@app.route("/users/<int:uid>/password", methods=["POST"])
def user_password(uid):
    gate = require_login()
    if gate: return gate

    conn = db()
    target = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not can_manage_user(target):
        conn.close()
        flash("لا تملك صلاحية تغيير كلمة مرور هذا الحساب")
        return redirect(url_for("users"))

    new_password = request.form.get("new_password", "").strip()
    if len(new_password) < 4:
        conn.close()
        flash("كلمة المرور يجب أن تكون 4 أحرف على الأقل")
        return redirect(url_for("users"))

    conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new_password), uid))
    conn.commit()
    conn.close()
    log_activity("تغيير كلمة مرور", f"تم تغيير كلمة مرور المستخدم رقم {uid}")
    flash("تم تغيير كلمة المرور")
    return redirect(url_for("users"))


@app.route("/expenses", methods=["GET","POST"])
def expenses():
    gate = require_login()
    if gate: return gate
    conn = db()
    if request.method == "POST":
        conn.execute("INSERT INTO expenses(title,amount,expense_date,notes,user_id) VALUES(?,?,?,?,?)",
                     (request.form["title"], float(request.form["amount"]), request.form["expense_date"], request.form.get("notes"), session["user_id"]))
        conn.commit()
    rows = conn.execute("SELECT * FROM expenses ORDER BY expense_date DESC, id DESC").fetchall()
    conn.close()
    return render_template("expenses.html", rows=rows, today=date.today().isoformat())





@app.route("/activity-logs")
def activity_logs():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    conn = db()
    rows = conn.execute("""SELECT l.*, u.full_name
                           FROM activity_logs l
                           LEFT JOIN users u ON u.id=l.user_id
                           ORDER BY l.id DESC LIMIT 300""").fetchall()
    conn.close()
    return render_template("activity_logs.html", rows=rows)




@app.route("/support")
def support():
    gate = require_login()
    if gate: return gate
    return render_template("support.html")


@app.route("/trash")
def trash():
    gate = require_login()
    if gate: return gate
    sup = super_admin_required()
    if sup: return sup
    return render_template("trash.html", rows=[])

@app.route("/currencies", methods=["GET","POST"])
def currencies():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    conn = db()
    conn.execute("""CREATE TABLE IF NOT EXISTS currencies (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT UNIQUE NOT NULL,
        name TEXT NOT NULL,
        rate_to_base REAL NOT NULL DEFAULT 1,
        is_base INTEGER DEFAULT 0,
        updated_at TEXT DEFAULT CURRENT_TIMESTAMP
    )""")
    defaults = [
        ("YER", "ريال يمني", 1, 1),
        ("SAR", "ريال سعودي", 140, 0),
        ("USD", "دولار أمريكي", 530, 0),
        ("AED", "درهم إماراتي", 145, 0)
    ]
    for code, name, rate, is_base in defaults:
        if not conn.execute("SELECT id FROM currencies WHERE code=?", (code,)).fetchone():
            conn.execute("INSERT INTO currencies(code,name,rate_to_base,is_base) VALUES(?,?,?,?)", (code, name, rate, is_base))
    conn.commit()

    if request.method == "POST":
        code = request.form["code"].upper().strip()
        name = request.form["name"].strip()
        rate = float(request.form.get("rate_to_base") or 1)
        is_base = int(request.form.get("is_base", 0))
        if is_base:
            conn.execute("UPDATE currencies SET is_base=0")
        conn.execute("""INSERT INTO currencies(code,name,rate_to_base,is_base,updated_at)
                        VALUES(?,?,?,?,CURRENT_TIMESTAMP)
                        ON CONFLICT(code) DO UPDATE SET
                        name=excluded.name,
                        rate_to_base=excluded.rate_to_base,
                        is_base=excluded.is_base,
                        updated_at=CURRENT_TIMESTAMP""",
                     (code, name, rate, is_base))
        conn.commit()
        flash("تم تحديث العملة وسعر الصرف")

    rows = conn.execute("SELECT * FROM currencies ORDER BY is_base DESC, code").fetchall()
    conn.close()
    return render_template("currencies.html", rows=rows)


@app.route("/tickets", methods=["GET","POST"])
def tickets():
    gate = require_login()
    if gate: return gate

    conn = db()
    clients_rows = conn.execute("SELECT * FROM clients ORDER BY full_name").fetchall()

    if request.method == "POST":
        client_id = request.form.get("client_id")
        if client_id == "new":
            cur = conn.execute("INSERT INTO clients(full_name,passport_number,phone,address,notes) VALUES(?,?,?,?,?)",
                               (request.form.get("new_full_name") or "عميل تذكرة", request.form.get("new_passport_number"),
                                request.form.get("new_phone"), "", "تم إنشاؤه من التذاكر"))
            client_id = cur.lastrowid

        image_path = ""
        file = request.files.get("ticket_image")
        if file and file.filename:
            ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else "jpg"
            filename = secure_filename(f"{uuid.uuid4().hex}.{ext}")
            save_path = Path(app.config["UPLOAD_FOLDER"]) / filename
            save_path.parent.mkdir(parents=True, exist_ok=True)
            file.save(save_path)
            image_path = "/" + str(save_path).replace("\\", "/")

        price = float(request.form.get("price") or 0)
        paid = float(request.form.get("paid_amount") or 0)
        remaining = price - paid

        conn.execute("""INSERT INTO tickets(
            client_id,ticket_type,route_from,route_to,travel_date,travel_time,seat_number,
            company_name,price,paid_amount,remaining_amount,image_path,notes,created_by
        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (client_id, request.form.get("ticket_type"), request.form.get("route_from"), request.form.get("route_to"),
         request.form.get("travel_date"), request.form.get("travel_time"), request.form.get("seat_number"),
         request.form.get("company_name"), price, paid, remaining, image_path, request.form.get("notes"), session["user_id"]))
        conn.commit()
        flash("تم حفظ التذكرة")
        return redirect(url_for("tickets"))

    rows = conn.execute("""SELECT tk.*, c.full_name, c.passport_number, c.phone
                           FROM tickets tk LEFT JOIN clients c ON c.id=tk.client_id
                           ORDER BY tk.id DESC""").fetchall()
    conn.close()
    return render_template("tickets.html", rows=rows, clients=clients_rows, today=date.today().isoformat())

@app.route("/annual-reports")
def annual_reports():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    year = request.args.get("year", str(date.today().year))
    conn = db()
    months = []
    for m in range(1, 13):
        ym = f"{year}-{m:02d}"
        tx_count = conn.execute("SELECT COUNT(*) c FROM transactions WHERE strftime('%Y-%m', created_at)=?", (ym,)).fetchone()["c"]
        income = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE strftime('%Y-%m', paid_at)=?", (ym,)).fetchone()["s"]
        expenses_sum = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE strftime('%Y-%m', expense_date)=?", (ym,)).fetchone()["s"]
        months.append({"month": ym, "tx_count": tx_count, "income": income, "expenses": expenses_sum, "profit": income-expenses_sum})
    years = conn.execute("SELECT DISTINCT strftime('%Y', created_at) y FROM transactions ORDER BY y DESC").fetchall()
    conn.close()
    return render_template("annual_reports.html", months=months, year=year, years=years)

@app.route("/whatsapp", methods=["GET","POST"])
def whatsapp():
    gate = require_login()
    if gate: return gate
    conn = db()

    if request.method == "POST":
        phone = request.form["phone"].strip()
        message = request.form["message"].strip()
        send_type = request.form["send_type"]
        scheduled_at = request.form.get("scheduled_at")

        delay_value = request.form.get("delay_value", "").strip()
        delay_unit = request.form.get("delay_unit", "minutes")
        if send_type == "scheduled" and delay_value:
            amount = int(delay_value)
            if delay_unit == "minutes":
                scheduled_at = (datetime.now() + timedelta(minutes=amount)).strftime("%Y-%m-%dT%H:%M")
            elif delay_unit == "hours":
                scheduled_at = (datetime.now() + timedelta(hours=amount)).strftime("%Y-%m-%dT%H:%M")
            elif delay_unit == "days":
                scheduled_at = (datetime.now() + timedelta(days=amount)).strftime("%Y-%m-%dT%H:%M")

        conn.execute("""INSERT INTO whatsapp_messages(phone,message,send_type,scheduled_at,status,created_by)
                        VALUES(?,?,?,?,?,?)""",
                     (phone, message, send_type, scheduled_at, "pending", session["user_id"]))
        conn.commit()

        if send_type == "manual":
            url = "https://wa.me/" + phone.replace("+","").replace(" ","") + "?text=" + quote(message)
            conn.close()
            return redirect(url)

        flash("تمت جدولة الرسالة وسيتم إرسالها تلقائياً عند الموعد إذا كان رابط API مفعلاً")
        conn.close()
        return redirect(url_for("whatsapp"))

    rows = conn.execute("SELECT * FROM whatsapp_messages ORDER BY id DESC").fetchall()
    settings = {r["key"]: r["value"] for r in conn.execute("SELECT * FROM settings").fetchall()}
    conn.close()
    return render_template("whatsapp.html", rows=rows, settings=settings)


@app.route("/whatsapp/settings", methods=["GET","POST"])
def whatsapp_settings():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    conn = db()
    if request.method == "POST":
        for key in ["whatsapp_api_url", "whatsapp_api_method", "whatsapp_api_headers", "whatsapp_enabled", "manager_notify_phone"]:
            conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, request.form.get(key, "")))
        conn.commit()
        flash("تم حفظ إعدادات واتساب")
    settings = {r["key"]: r["value"] for r in conn.execute("SELECT * FROM settings").fetchall()}
    conn.close()
    return render_template("whatsapp_settings.html", settings=settings)

@app.route("/notifications")
def notifications():
    gate = require_login()
    if gate: return gate
    conn = db()
    rows = conn.execute("SELECT * FROM notifications ORDER BY id DESC LIMIT 100").fetchall()
    conn.execute("UPDATE notifications SET is_read=1")
    conn.commit()
    conn.close()
    return render_template("notifications.html", rows=rows)

@app.route("/whatsapp/<int:msg_id>/send-now", methods=["POST"])
def whatsapp_send_now(msg_id):
    gate = require_login()
    if gate: return gate

    conn = db()
    msg = conn.execute("SELECT * FROM whatsapp_messages WHERE id=?", (msg_id,)).fetchone()
    conn.close()
    if not msg:
        flash("الرسالة غير موجودة")
        return redirect(url_for("whatsapp"))

    ok, response = send_whatsapp_api(msg["phone"], msg["message"])
    conn = db()
    if ok:
        conn.execute("UPDATE whatsapp_messages SET status='sent', sent_at=CURRENT_TIMESTAMP, provider_response=? WHERE id=?",
                     (response, msg_id))
        conn.execute("INSERT INTO notifications(title, body, type) VALUES(?,?,?)",
                     ("تم إرسال رسالة واتساب", f"تم إرسال رسالة إلى {msg['phone']}", "success"))
        flash("تم إرسال الرسالة")
    else:
        conn.execute("UPDATE whatsapp_messages SET status='failed', provider_response=? WHERE id=?",
                     (response, msg_id))
        conn.execute("INSERT INTO notifications(title, body, type) VALUES(?,?,?)",
                     ("فشل إرسال رسالة واتساب", f"الرقم: {msg['phone']} — السبب: {response}", "danger"))
        flash("فشل إرسال الرسالة، راجع التنبيهات أو إعدادات API")
    conn.commit()
    conn.close()
    return redirect(url_for("whatsapp"))

@app.route("/reports")
def reports():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm
    day = request.args.get("day", date.today().isoformat())
    conn = db()
    income = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE date(paid_at)=date(?)", (day,)).fetchone()["s"]
    expenses_sum = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE date(expense_date)=date(?)", (day,)).fetchone()["s"]
    txs = conn.execute("""SELECT t.*, c.full_name, c.passport_number, c.phone
                          FROM transactions t JOIN clients c ON c.id=t.client_id
                          WHERE COALESCE(t.is_deleted,0)=0 AND date(t.created_at)=date(?) ORDER BY t.id DESC""", (day,)).fetchall()
    conn.close()
    return render_template("reports.html", day=day, income=income, expenses_sum=expenses_sum, profit=income-expenses_sum, txs=txs)


@app.route("/reports/export")
def reports_export():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    period = request.args.get("period", "day")
    selected = request.args.get("date", date.today().isoformat())

    where = "date(t.created_at)=date(?)"
    params = [selected]
    title_period = f"تقرير يوم {selected}"

    if period == "month":
        where = "strftime('%Y-%m', t.created_at)=?"
        params = [selected[:7]]
        title_period = f"تقرير شهر {selected[:7]}"
    elif period == "year":
        where = "strftime('%Y', t.created_at)=?"
        params = [selected[:4]]
        title_period = f"تقرير سنة {selected[:4]}"

    conn = db()
    txs = conn.execute(f"""SELECT t.*, c.full_name, c.passport_number, c.phone, u.full_name employee
                           FROM transactions t
                           JOIN clients c ON c.id=t.client_id
                           LEFT JOIN users u ON u.id=t.employee_id
                           WHERE COALESCE(t.is_deleted,0)=0 AND {where}
                           ORDER BY t.created_at DESC""", params).fetchall()

    # Income and expenses periods
    if period == "day":
        income = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE date(paid_at)=date(?)", params).fetchone()["s"]
        expenses_sum = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE date(expense_date)=date(?)", params).fetchone()["s"]
    elif period == "month":
        income = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE strftime('%Y-%m', paid_at)=?", params).fetchone()["s"]
        expenses_sum = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE strftime('%Y-%m', expense_date)=?", params).fetchone()["s"]
    else:
        income = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM payments WHERE strftime('%Y', paid_at)=?", params).fetchone()["s"]
        expenses_sum = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM expenses WHERE strftime('%Y', expense_date)=?", params).fetchone()["s"]
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "التقرير"
    ws.sheet_view.rightToLeft = True

    ws.merge_cells("A1:N1")
    ws["A1"] = title_period
    ws["A1"].font = Font(size=16, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A3"] = "الإيرادات"
    ws["B3"] = income
    ws["C3"] = "المصروفات"
    ws["D3"] = expenses_sum
    ws["E3"] = "الصافي"
    ws["F3"] = income - expenses_sum

    headers = ["التاريخ", "اليوم", "الساعة", "العميل", "رقم الجواز", "الهاتف", "الخدمة", "التفويض", "رخصة قيادة", "تعديل مهنة", "الإجمالي", "المدفوع", "المتبقي", "الموظف"]
    ws.append([])
    ws.append(headers)

    for r in txs:
        try:
            dt = datetime.strptime(r["created_at"][:19], "%Y-%m-%d %H:%M:%S")
        except:
            dt = datetime.now()
        ws.append([
            r["created_at"][:10],
            arabic_day(dt),
            r["created_at"][11:19],
            r["full_name"],
            r["passport_number"],
            r["phone"],
            r["service_name"],
            r["delegated"],
            r["needs_license"],
            r["needs_profession_change"],
            r["total_amount"],
            r["paid_amount"],
            r["remaining_amount"],
            r["employee"] or "",
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="DDDDDD")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[5]:
        cell.fill = header_fill
        cell.font = header_font

    widths = [14,12,12,24,18,18,16,12,14,14,12,12,12,18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    filename = f"travel_agency_report_{period}_{selected}.xlsx"
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/settings", methods=["GET","POST"])
def settings():
    gate = require_login()
    if gate: return gate
    adm = admin_required()
    if adm: return adm
    conn = db()
    if request.method == "POST":
        for key in ["agency_name","default_alert_days"]:
            conn.execute("INSERT OR REPLACE INTO settings(key,value) VALUES(?,?)", (key, request.form.get(key, "")))
        conn.execute("UPDATE agencies SET name=? WHERE id=(SELECT id FROM agencies LIMIT 1)", (request.form.get("agency_name", "سفريات برو"),))
        conn.commit()
        flash("تم حفظ الإعدادات")
    rows = {r["key"]: r["value"] for r in conn.execute("SELECT * FROM settings").fetchall()}
    conn.close()
    return render_template("settings.html", rows=rows)






@app.route("/backup")
def backup():
    gate = require_login()
    if gate:
        return gate

    try:
        backup_folder = os.path.join(os.getcwd(), "backups")
        os.makedirs(backup_folder, exist_ok=True)

        if not os.path.exists(DB):
            init_db()

        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        db_name = "safaryat_pro_backup_" + stamp + ".db"
        zip_name = "safaryat_pro_backup_" + stamp + ".zip"

        db_copy_path = os.path.join(backup_folder, db_name)
        zip_path = os.path.join(backup_folder, zip_name)

        shutil.copy2(DB, db_copy_path)

        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
            z.write(db_copy_path, db_name)

        try:
            os.remove(db_copy_path)
        except Exception:
            pass

        return send_from_directory(
            directory=os.path.abspath(backup_folder),
            path=zip_name,
            as_attachment=True,
            mimetype="application/zip"
        )

    except Exception as e:
        flash("خطأ أثناء إنشاء النسخة الاحتياطية: " + str(e))
        return redirect(url_for("dashboard"))


if __name__ == "__main__":
    init_db()
    ensure_database_upgrade()
    start_scheduler_once()
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
